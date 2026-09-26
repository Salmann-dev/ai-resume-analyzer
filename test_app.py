import io
import os
import unittest
from unittest.mock import patch, MagicMock
from docx import Document
from app import app, allowed_file, analyze_resume_with_gemini

class ResumeAnalyzerTestCase(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        app.config["WTF_CSRF_ENABLED"] = False
        self.client = app.test_client()

    def test_allowed_file(self):
        self.assertTrue(allowed_file("resume.pdf"))
        self.assertTrue(allowed_file("resume.docx"))
        self.assertTrue(allowed_file("RESUME.PDF"))
        self.assertFalse(allowed_file("resume.txt"))
        self.assertFalse(allowed_file("script.py"))
        self.assertFalse(allowed_file("no_extension"))

    def test_homepage_loads(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"AI Resume", response.data)
        self.assertIn(b"Analyzer", response.data)
        self.assertIn(b'name="resume"', response.data)
        self.assertIn(b'name="job_description"', response.data)

    def test_analyze_no_file(self):
        response = self.client.post("/analyze", data={}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Please upload a resume file", response.data)

    def test_analyze_invalid_extension(self):
        data = {
            "resume": (io.BytesIO(b"dummy text"), "resume.txt"),
            "job_description": "Software Engineer"
        }
        response = self.client.post("/analyze", data=data, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Invalid file format", response.data)

    @patch("app.genai.Client")
    @patch.dict("os.environ", {"GEMINI_API_KEY": "fake_gemini_key"})
    def test_analyze_resume_with_gemini_parser(self, mock_genai_client_class):
        mock_client = MagicMock()
        mock_genai_client_class.return_value = mock_client

        mock_response = MagicMock()
        mock_response.text = '''```json
{
  "match_score": 88,
  "strengths": ["Python expertise", "Strong system design"],
  "weaknesses": ["Limited AWS cloud experience"],
  "missing_keywords": ["Docker", "Kubernetes"],
  "summary": "Excellent software engineering background with proven delivery experience."
}
```'''
        mock_client.models.generate_content.return_value = mock_response

        result = analyze_resume_with_gemini("Software engineer with 5 years experience in Python.", "Python Dev")
        self.assertEqual(result["match_score"], 88)
        self.assertEqual(len(result["strengths"]), 2)
        self.assertEqual(len(result["weaknesses"]), 1)
        self.assertIn("Docker", result["missing_keywords"])
        self.assertIn("Excellent", result["summary"])
        mock_client.models.generate_content.assert_called_once()
        args, kwargs = mock_client.models.generate_content.call_args
        self.assertEqual(kwargs.get("model"), "gemini-3.8-flash")

    @patch("app.analyze_resume_with_gemini")
    def test_full_analyze_route_with_docx(self, mock_analyzer):
        mock_analyzer.return_value = {
            "match_score": 92,
            "strengths": ["Full stack development proficiency", "Deep Python skills"],
            "weaknesses": ["Needs more public portfolio projects"],
            "missing_keywords": ["GraphQL", "Next.js"],
            "summary": "Outstanding candidate profile strongly matched to the engineering position."
        }

        doc = Document()
        doc.add_heading("Alex Morgan - Senior Backend Engineer", 0)
        doc.add_paragraph("Summary: 7+ years developing distributed backend systems in Python and Go.")
        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        doc_stream.seek(0)

        data = {
            "resume": (doc_stream, "alex_morgan_resume.docx"),
            "job_description": "Senior Backend Engineer with Python experience"
        }

        response = self.client.post("/analyze", data=data, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"92", response.data)
        self.assertIn(b"/100", response.data)
        self.assertIn(b"Full stack development proficiency", response.data)
        self.assertIn(b"Needs more public portfolio projects", response.data)
        self.assertIn(b"GraphQL", response.data)
        self.assertIn(b"Outstanding candidate profile", response.data)

if __name__ == "__main__":
    unittest.main()
